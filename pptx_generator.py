from io import BytesIO
from typing import List

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt
from lxml import etree

from text_parser import Deck, SlideContent

THEMES = {
    "light": {
        "bg": RGBColor(255, 255, 255),
        "title": RGBColor(33, 47, 61),
        "text": RGBColor(52, 73, 94),
        "accent": RGBColor(41, 128, 185),
        "muted": RGBColor(127, 140, 141),
        "panel": RGBColor(236, 243, 249),
    },
    "dark": {
        "bg": RGBColor(28, 40, 51),
        "title": RGBColor(236, 240, 241),
        "text": RGBColor(202, 208, 212),
        "accent": RGBColor(93, 173, 226),
        "muted": RGBColor(149, 165, 166),
        "panel": RGBColor(40, 57, 72),
    },
}

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
FONT = "Calibri"


def _set_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def _rect(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def _textbox(slide, left, top, width, height, text, size, color, bold=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def _set_bullet(paragraph, char="•"):
    """Настоящий маркер списка с висячим отступом (вместо символа в тексте)."""
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(0.4))))
    pPr.set("indent", str(-int(Inches(0.4))))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    bu = etree.SubElement(pPr, qn("a:buChar"))
    bu.set("char", char)


def _font_size_for(total_chars: int) -> int:
    if total_chars <= 250:
        return 28
    if total_chars <= 450:
        return 24
    if total_chars <= 700:
        return 20
    return 18


def _split_slide(slide: SlideContent, max_items: int, max_chars: int) -> List[SlideContent]:
    """Делит слишком длинный слайд на несколько, чтобы текст не вылезал."""
    items = [("b", b) for b in slide.bullets] + [("p", p) for p in slide.paragraphs]
    if not items:
        return [slide]

    parts, current, chars = [], [], 0
    for kind, text in items:
        if current and (len(current) >= max_items or chars + len(text) > max_chars):
            parts.append(current)
            current, chars = [], 0
        current.append((kind, text))
        chars += len(text)
    if current:
        parts.append(current)

    result = []
    for i, part in enumerate(parts, start=1):
        title = slide.title if len(parts) == 1 else f"{slide.title} ({i}/{len(parts)})"
        result.append(SlideContent(
            title=title,
            bullets=[t for k, t in part if k == "b"],
            paragraphs=[t for k, t in part if k == "p"],
        ))
    return result


def _add_footer(slide, theme, number, total):
    _textbox(slide, Inches(11.4), Inches(6.85), Inches(1.5), Inches(0.4),
             f"{number} / {total}", 12, theme["muted"], align=PP_ALIGN.RIGHT)


def build_presentation(deck: Deck, author: str = "", theme_name: str = "light",
                       max_items_per_slide: int = 6) -> BytesIO:
    """Собирает готовую презентацию из текстовой структуры (без какого-либо кода в слайдах)."""
    theme = THEMES[theme_name]
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    content_slides: List[SlideContent] = []
    for s in deck.slides:
        content_slides.extend(_split_slide(s, max_items_per_slide, 600))

    show_agenda = len(deck.slides) >= 3
    total = 1 + (1 if show_agenda else 0) + len(content_slides) + 1
    number = 0

    # 1. Титульный слайд
    slide = prs.slides.add_slide(blank)
    _set_bg(slide, theme["bg"])
    _rect(slide, 0, 0, Inches(0.5), SLIDE_H, theme["accent"])
    _textbox(slide, Inches(1.4), Inches(2.3), Inches(10.8), Inches(1.9), deck.title,
             48, theme["title"], bold=True, anchor=MSO_ANCHOR.BOTTOM)
    _rect(slide, Inches(1.5), Inches(4.4), Inches(2.2), Inches(0.08), theme["accent"])
    if deck.subtitle:
        _textbox(slide, Inches(1.4), Inches(4.7), Inches(10.8), Inches(1.2), deck.subtitle,
                 24, theme["muted"])
    if author:
        _textbox(slide, Inches(1.4), Inches(6.3), Inches(10.8), Inches(0.5), author,
                 18, theme["muted"])
    number += 1

    # 2. Содержание
    if show_agenda:
        slide = prs.slides.add_slide(blank)
        _set_bg(slide, theme["bg"])
        _rect(slide, 0, 0, SLIDE_W, Inches(1.3), theme["panel"])
        _rect(slide, 0, 0, Inches(0.25), Inches(1.3), theme["accent"])
        _textbox(slide, Inches(0.8), Inches(0.2), Inches(11.7), Inches(0.9), "Содержание",
                 36, theme["title"], bold=True, anchor=MSO_ANCHOR.MIDDLE)
        box = slide.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(4.8))
        tf = box.text_frame
        tf.word_wrap = True
        for i, s in enumerate(deck.slides[:10]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_after = Pt(10)
            run = p.add_run()
            run.text = f"{i + 1}.  {s.title}"
            run.font.name = FONT
            run.font.size = Pt(24 if len(deck.slides) <= 7 else 20)
            run.font.color.rgb = theme["text"]
        number += 1
        _add_footer(slide, theme, number, total)

    # 3. Основные слайды
    for s in content_slides:
        slide = prs.slides.add_slide(blank)
        _set_bg(slide, theme["bg"])
        _rect(slide, 0, 0, SLIDE_W, Inches(1.3), theme["panel"])
        _rect(slide, 0, 0, Inches(0.25), Inches(1.3), theme["accent"])
        _textbox(slide, Inches(0.8), Inches(0.2), Inches(11.7), Inches(0.9), s.title,
                 34 if len(s.title) <= 45 else 28, theme["title"], bold=True,
                 anchor=MSO_ANCHOR.MIDDLE)

        total_chars = sum(len(t) for t in s.bullets + s.paragraphs)
        size = _font_size_for(total_chars)

        box = slide.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(4.8))
        tf = box.text_frame
        tf.word_wrap = True
        first = True
        for text in s.paragraphs:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.space_after = Pt(14)
            run = p.add_run()
            run.text = text
            run.font.name = FONT
            run.font.size = Pt(size)
            run.font.color.rgb = theme["text"]
        for text in s.bullets:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.space_after = Pt(12)
            _set_bullet(p)
            run = p.add_run()
            run.text = text
            run.font.name = FONT
            run.font.size = Pt(size)
            run.font.color.rgb = theme["text"]
        number += 1
        _add_footer(slide, theme, number, total)

    # 4. Финальный слайд
    slide = prs.slides.add_slide(blank)
    _set_bg(slide, theme["bg"])
    _rect(slide, 0, 0, Inches(0.5), SLIDE_H, theme["accent"])
    _textbox(slide, Inches(1.4), Inches(2.8), Inches(10.8), Inches(1.4), "Спасибо за внимание!",
             48, theme["title"], bold=True, anchor=MSO_ANCHOR.MIDDLE)
    _rect(slide, Inches(1.5), Inches(4.3), Inches(2.2), Inches(0.08), theme["accent"])
    if author:
        _textbox(slide, Inches(1.4), Inches(4.6), Inches(10.8), Inches(0.6), author,
                 20, theme["muted"])

    prs.core_properties.title = deck.title
    if author:
        prs.core_properties.author = author

    buffer = BytesIO()
    prs.save(buffer)
    buffer.seek(0)
    return buffer
